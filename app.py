import io
import re
import time
import gzip
import json
import hashlib
import ipaddress
from collections import deque
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from urllib.parse import urljoin, urlparse, urldefrag, quote, unquote
from email.utils import parsedate_to_datetime
from urllib.robotparser import RobotFileParser
import socket
import threading

import pandas as pd
import requests
import streamlit as st
from bs4 import BeautifulSoup
from openpyxl.styles import Font

st.set_page_config(page_title='掲載URL収集ツール', layout='wide')
USER_AGENT = 'ListingAuditBot/1.0 (respectful site audit; contact your site administrator)'
TIMEOUT = 12
MAX_BYTES = 2_000_000
HEADERS = {'User-Agent': USER_AGENT, 'Accept': 'text/html,application/xml,text/xml;q=0.9,*/*;q=0.5'}


def host_of(url):
    try:
        return (urlparse(url).hostname or '').lower().rstrip('.')
    except ValueError:
        return ''


def normalize_host(raw):
    raw = str(raw).strip()
    if not raw or raw.lower() in ('nan', 'none'): return ''
    if not re.match(r'^https?://', raw, re.I): raw = 'https://' + raw
    h = host_of(raw)
    if h.startswith('www.'): h = h[4:]
    if not h or '.' not in h or h in {'google.co.jp', 'google.com'} or h.endswith('.google.com') or h.endswith('.google.co.jp'):
        return ''
    return h


def in_scope(host, domain):
    return host == domain or host.endswith('.' + domain)


def public_url(url, domain=None):
    try:
        p = urlparse(url)
        if p.scheme not in ('http', 'https') or not p.hostname or p.username or p.password: return False
        h = p.hostname.lower()
        if domain and not in_scope(h, domain): return False
        # Block private/reserved IPs and names that resolve to them (SSRF protection).
        addresses = socket.getaddrinfo(h, p.port or (443 if p.scheme == 'https' else 80), type=socket.SOCK_STREAM)
        if not addresses: return False
        return all(ipaddress.ip_address(a[4][0]).is_global for a in addresses)
    except (OSError, ValueError):
        return False


# Each worker owns its diagnostic records; crawling decisions remain unchanged.
_diag_local = threading.local()

def log_request(url, status, detail='', elapsed=0.0):
    records = getattr(_diag_local, 'records', None)
    if records is not None:
        records.append({'ドメイン': getattr(_diag_local, 'domain', ''),
                        '取得URL': url, 'HTTPステータス': status,
                        '詳細': str(detail)[:300], '所要秒': round(elapsed, 2)})

def log_page_diagnostic(requested, final_url, title, text_length, text_match, html_match):
    records = getattr(_diag_local, 'page_records', None)
    if records is not None:
        records.append({'ドメイン': getattr(_diag_local, 'domain', ''),
                        '確認URL': requested, '最終URL': final_url, 'ページタイトル': title,
                        '抽出本文文字数': text_length, '本文キーワード一致': '○' if text_match else '×',
                        'HTMLキーワード一致': '○' if html_match else '×',
                        '判定メモ': ('本文一致' if text_match else
                                  'HTML内のみ一致（非表示要素等の可能性）' if html_match else
                                  'HTML・抽出本文とも一致なし')})


def safe_get(url, domain, session):
    start = time.monotonic()
    if not public_url(url, domain):
        log_request(url, 'URL検証不可', '公開URLとして検証できないかDNS解決に失敗', time.monotonic()-start)
        return None
    try:
        r = session.get(url, headers=HEADERS, timeout=TIMEOUT, allow_redirects=False, stream=True)
        if r.status_code in (301, 302, 303, 307, 308):
            location = r.headers.get('Location')
            r.close()
            if not location:
                log_request(url, r.status_code, 'Locationヘッダーなし', time.monotonic()-start)
                return None
            target = urljoin(url, location)
            if not public_url(target, domain):
                log_request(url, r.status_code, 'リダイレクト先が許可範囲外または検証不可', time.monotonic()-start)
                return None
            r = session.get(target, headers=HEADERS, timeout=TIMEOUT, allow_redirects=False, stream=True)
        # 一時的なサーバーエラーだけ控えめに再試行。403や429は再試行しない。
        if r.status_code in (500, 502, 503, 504):
            first_status = r.status_code
            r.close()
            log_request(url, first_status, '一時的なサーバーエラー・再試行待ち', time.monotonic()-start)
            time.sleep(2)
            r = session.get(url, headers=HEADERS, timeout=TIMEOUT, allow_redirects=False, stream=True)
        if r.status_code != 200:
            log_request(url, r.status_code, 'HTTPエラー', time.monotonic()-start)
            r.close()
            return None
        typ = r.headers.get('Content-Type', '').lower()
        if not any(t in typ for t in ('html', 'xml', 'text/plain')):
            log_request(url, 200, '対象外Content-Type: '+typ, time.monotonic()-start)
            r.close()
            return None
        raw = bytearray()
        for chunk in r.iter_content(65536):
            raw.extend(chunk)
            if len(raw) > MAX_BYTES:
                log_request(url, 200, 'サイズ上限超過', time.monotonic()-start)
                r.close()
                return None
        r._content = bytes(raw)
        log_request(url, 200, '取得成功', time.monotonic()-start)
        return r
    except requests.Timeout as exc:
        log_request(url, 'タイムアウト', type(exc).__name__, time.monotonic()-start)
        return None
    except requests.RequestException as exc:
        log_request(url, '通信エラー', type(exc).__name__+': '+str(exc), time.monotonic()-start)
        return None


def allowed(url, robots):
    try: return robots.can_fetch(USER_AGENT, url)
    except Exception: return False


def load_robots(base, domain, session):
    robots = RobotFileParser()
    url = base + '/robots.txt'
    response = safe_get(url, domain, session)
    if response is None:
        # Fail closed if robots rules cannot be retrieved.
        return None, []
    lines = response.text.splitlines()
    robots.parse(lines)
    maps = [line.split(':', 1)[1].strip() for line in lines if line.lower().startswith('sitemap:')]
    return robots, maps


def sitemap_urls(base, domain, session, robots, extra_maps, max_pages):
    queue = deque(dict.fromkeys(extra_maps + [base + '/sitemap.xml']))
    visited = set()
    results = []
    while queue and len(visited) < 25 and len(results) < max_pages:
        url = queue.popleft()
        if url in visited or not public_url(url, domain) or not allowed(url, robots): continue
        visited.add(url)
        r = safe_get(url, domain, session)
        if r is None: continue
        try:
            soup = BeautifulSoup(r.content, 'xml')
            for loc in soup.find_all('loc'):
                link = loc.get_text(strip=True)
                if not link or not in_scope(host_of(link), domain): continue
                if loc.parent and loc.parent.name.lower() == 'sitemap':
                    if link not in visited: queue.append(link)
                elif loc.parent and loc.parent.name.lower() == 'url':
                    results.append(link)
                    if len(results) >= max_pages: break
        except Exception:
            continue
    return list(dict.fromkeys(results))


def extract_site_name(soup, domain):
    """Best-effort site name, not article title; never invent a name."""
    for attrs in ({'property': 'og:site_name'}, {'name': 'application-name'},
                  {'name': 'apple-mobile-web-app-title'}):
        tag = soup.find('meta', attrs=attrs)
        if tag and tag.get('content', '').strip():
            return tag['content'].strip()[:120]
    if soup.title:
        title = soup.title.get_text(' ', strip=True)
        # Common article-title | site-name and article-title - site-name formats.
        for sep in (' | ', '｜', ' - ', ' – ', ' — ', '｜'):
            if sep in title:
                candidate = title.rsplit(sep, 1)[-1].strip()
                if candidate and len(candidate) <= 80:
                    return candidate
    return ''


def check_page(url, domain, session, robots, keyword, link_patterns):
    if not allowed(url, robots): return None, []
    r = safe_get(url, domain, session)
    if r is None:
        return None, []
    if 'html' not in r.headers.get('Content-Type', '').lower():
        log_request(url, 'HTML対象外', '本文判定を行わずスキップ')
        return None, []
    soup = BeautifulSoup(r.content, 'html.parser')
    anchors = [(urljoin(r.url, a.get('href', '')), a.get_text(' ', strip=True)) for a in soup.find_all('a', href=True)]
    links = sorted(set(href for href, _ in anchors))  # HTMLのリンク順による探索差を減らす
    # パターン未指定時は「候補」のみ検出。ASPの中継先は検証しない。
    if link_patterns:
        matched_links = [href for href, _ in anchors if any(p.casefold() in href.casefold() for p in link_patterns)]
        link_status = 'パターン一致' if matched_links else 'なし'
    else:
        matched_links = [href for href, label in anchors if keyword and
                         (keyword.casefold() in label.casefold() or keyword.casefold() in href.casefold())]
        link_status = '候補・要確認' if matched_links else 'なし'
    for tag in soup(['script', 'style', 'noscript']): tag.decompose()
    text = soup.get_text(' ', strip=True)
    text_match = keyword.casefold() in text.casefold() if keyword else False
    # Keep raw HTML separate: helps distinguish dynamic-rendered text from text extraction.
    html_match = keyword.casefold() in r.text.casefold() if keyword else False
    title_text = soup.title.get_text(' ', strip=True)[:200] if soup.title else ''
    log_page_diagnostic(url, r.url, title_text, len(text), text_match, html_match)
    return {'url': r.url, 'text_match': text_match, 'link_match': bool(matched_links),
            'link_status': link_status,
            'matched_links': ' | '.join(dict.fromkeys(matched_links))[:3000],
            'title': title_text,
            'site_name': extract_site_name(soup, domain)}, links


def read_urls(upload):
    name = upload.name.lower()
    if name.endswith('.csv'):
        raw = upload.getvalue()
        for enc in ('utf-8-sig', 'cp932'):
            try: return pd.read_csv(io.BytesIO(raw), dtype=str, encoding=enc)
            except UnicodeDecodeError: pass
        raise ValueError('CSVの文字コードを読み取れません')
    return pd.read_excel(upload, dtype=str)


def excel_bytes(details, summary, errors, history, diagnostics, page_diagnostics):
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine='openpyxl') as writer:
        for name, frame in [('掲載URL一覧', details), ('ドメイン別集計', summary), ('エラー一覧', errors), ('検索履歴', history), ('取得診断', diagnostics), ('ページ判定診断', page_diagnostics)]:
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = 'A2'
            for row in ws.iter_rows():
                for cell in row:
                    cell.font = Font(name='Meiryo UI', size=10, bold=(cell.row == 1))
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                letter = col[0].column_letter
                ws.column_dimensions[letter].width = min(65, max(15, max(len(str(cell.value or '')) for cell in list(col)[:100]) + 2))
    return out.getvalue()


st.title('🔎 掲載URL収集ツール')
st.markdown('''<style>
.stButton > button[kind="primary"] {background: linear-gradient(105deg,#2563eb,#7c3aed); color:white; border:0; border-radius:14px; min-height:3.4rem; font-size:1.08rem; font-weight:750; box-shadow:0 5px 15px rgba(67,56,202,.22);}
.stButton > button[kind="primary"]:hover {background:linear-gradient(105deg,#1d4ed8,#6d28d9); color:white; transform:translateY(-1px);}
</style>''', unsafe_allow_html=True)
st.caption('サイトマップ・サイト内リンクを検査。掲載判定と未確認の理由を診断します。')
st.subheader('URL一覧の入力')
input_method = st.radio('入力方法', ['ファイルをアップロード', 'テキストを直接貼り付け'], horizontal=True)
file = None
pasted = ''
if input_method == 'ファイルをアップロード':
    file = st.file_uploader('URL一覧（Excel / CSV）', type=['xlsx', 'csv'])
else:
    pasted = st.text_area('URL一覧を貼り付け（1行に1件）', height=180,
                          placeholder='https://example.com/page1\nhttps://www.example.org/page2')
    st.caption('ExcelのURL列をコピーして、そのまま貼り付けてもOK。空行は無視します。')
keyword = st.text_input('検索キーワード', value='楽天銀行スーパーローン')
mode = st.radio('掲載URLの検出条件', ['商品名の掲載', '送客リンクの掲載', '両方を検出'], index=2)
st.caption('送客リンクはリンクテキストまたはURLに検索キーワードが含まれるものを候補として抽出します（要確認）。')
max_pages = st.number_input('ドメインあたり最大確認ページ数', 10, 1000, 10, step=10)
delay = st.number_input('同一サイトへのアクセス間隔（秒）', 0.5, 10.0, 1.0, step=0.5)

raw_urls = []
if input_method == 'ファイルをアップロード' and file is not None:
    try:
        data = read_urls(file)
        columns = list(data.columns)
        url_col = st.selectbox('URLが入っている列', columns,
                               index=next((i for i,c in enumerate(columns) if 'url' in str(c).lower()), 0))
        raw_urls = data[url_col].dropna().tolist()
    except Exception as exc:
        st.error(f'入力ファイルの読込エラー: {exc}')
elif input_method == 'テキストを直接貼り付け' and pasted.strip():
    raw_urls = [line.strip() for line in pasted.splitlines() if line.strip()]

domains = list(dict.fromkeys(filter(None, (normalize_host(v) for v in raw_urls))))
if raw_urls:
    st.info(f'入力URL数: {len(raw_urls):,} 件 / 抽出ドメイン数: {len(domains):,} 件')
    with st.expander('ドメイン一覧を確認'):
        st.dataframe(pd.DataFrame({'サイト名': [st.session_state.get('site_names', {}).get(d, '') for d in domains], 'ドメイン': domains}), hide_index=True)
    domain_excel = io.BytesIO()
    with pd.ExcelWriter(domain_excel, engine='openpyxl') as writer:
        pd.DataFrame({'サイト名': [st.session_state.get('site_names', {}).get(d, '') for d in domains], 'ドメイン': domains}).to_excel(writer, index=False, sheet_name='ドメイン一覧')
        ws = writer.sheets['ドメイン一覧']
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 45
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name='Meiryo UI', size=10, bold=(cell.row == 1))
    st.download_button('📥 ドメイン一覧をダウンロード', domain_excel.getvalue(),
                       file_name='ドメイン一覧.xlsx',
                       mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                       on_click='ignore')

# v1.5: 公開インデックス照会は停止。取得判定の診断を優先する。

def url_priority(url, keyword):
    """Higher scores first, deterministic; only affects URLs after explicit inputs."""
    p = urlparse(url)
    path = unquote(p.path + '?' + p.query).casefold()
    compact = re.sub(r'\s+', '', keyword).casefold()
    score = 0
    if compact and compact in re.sub(r'\s+', '', path): score += 100
    for token, weight in [('rakuten', 24), ('楽天', 24), ('bank', 12), ('銀行', 12),
                          ('loan', 24), ('ローン', 24), ('super', 8), ('スーパー', 8),
                          ('card', 6), ('finance', 6), ('借入', 10), ('審査', 6)]:
        if token in path: score += weight
    return (-score, len(p.path), url)


def collect_domain(domain, keyword, mode, max_pages, delay, settings_id, seed_urls):
    details, errors = [], []
    _diag_local.records = []
    _diag_local.page_records = []
    _diag_local.domain = domain
    site_name = ''
    now = datetime.now(timezone.utc).isoformat(timespec='seconds')
    visited, checked, succeeded = set(), 0, 0
    skipped_robots = 0
    selected_urls = []
    session = requests.Session()  # ワーカー間でSessionを共有しない
    try:
        base = 'https://' + domain
        robots, maps = load_robots(base, domain, session)
        if robots is None:
            raise RuntimeError('robots.txtを確認できないため巡回をスキップしました')
        crawl_delay = robots.crawl_delay(USER_AGENT)
        actual_delay = max(float(delay), float(crawl_delay or 0))
        # サイトマップは候補収集専用。実際のページ確認は max_pages 以内。
        # 入力された掲載URLを最優先で検査。旧版で拾えたURLの取りこぼしを防ぐ。
        sitemap_candidates = sitemap_urls(base, domain, session, robots, maps, int(max_pages) * 3)
        candidates = list(seed_urls) + sorted(sitemap_candidates, key=lambda u: url_priority(u, keyword))
        candidates.append(base + '/')
        queued = set()
        frontier = []
        def add(url):
            url = urldefrag(url)[0]
            if url not in queued and url not in visited and in_scope(host_of(url), domain):
                queued.add(url)
                frontier.append(url)
        for url in candidates: add(url)
        seed_count = len(frontier)
        last_request_at = 0.0
        while frontier and checked < max_pages:
            url = frontier.pop(0)
            if url in visited: continue
            if not allowed(url, robots):
                skipped_robots += 1
                log_request(url, '巡回禁止', 'robots.txtのルールにより確認対象外')
                continue
            visited.add(url)
            selected_urls.append(url)
            # 同一ドメインのページ確認間隔。異なるドメインは並列に進む。
            remaining = actual_delay - (time.monotonic() - last_request_at)
            if remaining > 0: time.sleep(remaining)
            last_request_at = time.monotonic()
            checked += 1
            page, links = check_page(url, domain, session, robots, keyword, [])
            if page:
                if not site_name and page.get('site_name'):
                    site_name = page['site_name']
                succeeded += 1
                match = ((mode == '商品名の掲載' and page['text_match']) or
                         (mode == '送客リンクの掲載' and page['link_match']) or
                         (mode == '両方を検出' and (page['text_match'] or page['link_match'])))
                if match:
                    details.append({'サイト名': site_name or page.get('site_name', ''), 'ドメイン': domain, 'URL': page['url'], 'タイトル': page['title'],
                                    '商品名掲載': '○' if page['text_match'] else '×',
                                    '送客リンク掲載': page['link_status'],
                                    '一致したリンク': page['matched_links'], '確認日時UTC': now,
                                    '設定ID': settings_id})
                for link in sorted(links, key=lambda u: url_priority(u, keyword)): add(link)
        failures = [x for x in _diag_local.records if x['HTTPステータス'] in ('通信エラー','タイムアウト','URL検証不可') or (isinstance(x['HTTPステータス'], int) and x['HTTPステータス'] >= 400)]
        if checked == 0:
            state = '巡回制限・未確認' if skipped_robots else '確認ページなし'
        elif succeeded == 0:
            state = '取得失敗・未確認'
        elif failures or succeeded < checked:
            state = '一部取得失敗'
        elif frontier and checked >= max_pages:
            state = '確認上限到達'
        else:
            state = '確認完了（確認範囲内）'
    except Exception as exc:
        errors.append({'ドメイン': domain, 'URL': '', 'エラー': str(exc)})
        state = '巡回制限・未確認' if 'robots.txt' in str(exc) else '取得失敗・未確認'
    finally:
        session.close()
    history = {'サイト名': site_name, 'ドメイン': domain, '設定ID': settings_id, '状態': state,
               '確認ページ数': checked, '取得成功ページ数': succeeded, '候補URL数': len(locals().get('frontier', [])) + len(visited), '検出URL数': len(details), '確認日時UTC': now, '確認対象URL': ' | '.join(selected_urls)[:30000], 'robots除外数': skipped_robots, '判定診断ページ数': len(_diag_local.page_records)}
    records = _diag_local.records
    page_records = _diag_local.page_records
    _diag_local.records = None
    _diag_local.page_records = None
    return details, errors, history, records, page_records


st.caption('ドメイン間は5件並列。入力URLを最優先、サイトマップとサイト内リンクは関連語を含むURLを優先します。')
workers = 5  # 同時処理ドメイン数（UIには表示しない）
if st.button('🚀 掲載URLを収集する', type='primary', disabled=not domains, use_container_width=True):
    if not keyword.strip(): st.error('検索キーワードを入力してください。'); st.stop()
    settings_id = hashlib.sha256(json.dumps({'keyword': keyword, 'mode': mode}, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:16]
    st.session_state.pop('last_result', None)
    seed_map = {d: [] for d in domains}
    for raw in raw_urls:
        d = normalize_host(raw)
        if d in seed_map:
            candidate = str(raw).strip()
            if not re.match(r'^https?://', candidate, re.I): candidate = 'https://' + candidate
            if public_url(candidate, d) and candidate not in seed_map[d]: seed_map[d].append(candidate)
    progress = st.progress(0)
    status = st.empty()
    details, errors, history, diagnostics, page_diagnostics = [], [], [], [], []
    with ThreadPoolExecutor(max_workers=int(workers)) as pool:
        futures = {pool.submit(collect_domain, domain, keyword.strip(), mode, int(max_pages), float(delay), settings_id, seed_map[domain]): domain for domain in domains}
        for idx, future in enumerate(as_completed(futures), start=1):
            domain = futures[future]
            try:
                d, e, h, diag, pages = future.result()
                details.extend(d); errors.extend(e); history.append(h); diagnostics.extend(diag); page_diagnostics.extend(pages)
            except Exception as exc:
                errors.append({'ドメイン': domain, 'URL': '', 'エラー': str(exc)})
            progress.progress(idx / len(domains))
            status.write(f'完了 {idx}/{len(domains)} ドメイン（直近: {domain}）')
    df = pd.DataFrame(details, columns=['サイト名','ドメイン','URL','タイトル','商品名掲載','送客リンク掲載','一致したリンク','確認日時UTC','設定ID'])
    hdf = pd.DataFrame(history, columns=['サイト名','ドメイン','設定ID','状態','確認ページ数','取得成功ページ数','候補URL数','検出URL数','確認日時UTC','確認対象URL','robots除外数','判定診断ページ数'])
    summary = pd.DataFrame({'ドメイン': domains}).merge(df.groupby('ドメイン').size().rename('検出URL件数').reset_index(), how='left', on='ドメイン')
    name_map = dict(zip(hdf['ドメイン'], hdf['サイト名']))
    st.session_state['site_names'] = name_map
    summary.insert(0, 'サイト名', summary['ドメイン'].map(name_map).fillna(''))
    if not df.empty:
        df['サイト名'] = df['ドメイン'].map(name_map).fillna(df['サイト名'])
    summary['検出URL件数'] = summary['検出URL件数'].fillna(0).astype(int)
    summary = summary.merge(hdf[['ドメイン','状態','確認ページ数','取得成功ページ数','判定診断ページ数']], on='ドメイン', how='left')
    edf = pd.DataFrame(errors, columns=['ドメイン','URL','エラー'])
    diagdf = pd.DataFrame(diagnostics, columns=['ドメイン','取得URL','HTTPステータス','詳細','所要秒'])
    pagedf = pd.DataFrame(page_diagnostics, columns=['ドメイン','確認URL','最終URL','ページタイトル','抽出本文文字数','本文キーワード一致','HTMLキーワード一致','判定メモ'])
    st.session_state['last_result'] = {
        'df': df, 'history': hdf, 'errors': edf, 'diagnostics': diagdf, 'page_diagnostics': pagedf, 'excel': excel_bytes(df, summary, edf, hdf, diagdf, pagedf),
        'filename': f'掲載URL収集結果_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx',
        'detected': len(details), 'error_count': len(errors),
    }
    st.balloons()

# ダウンロードによる再実行でも結果を消さず、次の収集開始まで保持する。
if 'last_result' in st.session_state:
    result = st.session_state['last_result']
    # Streamlit Community Cloud may preserve session_state across code redeploys.
    # Older results have no diagnostics key; migrate them without crashing.
    if 'diagnostics' not in result:
        result['diagnostics'] = pd.DataFrame(columns=['ドメイン','取得URL','HTTPステータス','詳細','所要秒'])
        st.info('前バージョンの収集結果を表示しています。取得診断を記録するには「収集開始」をもう一度押してください。')
    if 'page_diagnostics' not in result:
        result['page_diagnostics'] = pd.DataFrame(columns=['ドメイン','確認URL','最終URL','ページタイトル','抽出本文文字数','本文キーワード一致','HTMLキーワード一致','判定メモ'])
    if 'errors' not in result:
        result['errors'] = pd.DataFrame(columns=['ドメイン','URL','エラー'])
    st.success(f"処理終了：今回の検出 {result['detected']} 件 / エラー {result['error_count']} 件")
    st.subheader('検索結果')
    st.caption('※ 入力URLを優先して確認します。検出0件は掲載なしを意味しません。巡回制限・取得失敗・確認上限を診断欄で区別します。')
    st.dataframe(result['df'], hide_index=True, use_container_width=True)
    with st.expander('ドメイン別の確認状況（診断）'):
        st.dataframe(result['history'], hide_index=True, use_container_width=True)
    with st.expander('取得診断（HTTPステータス・通信エラー）'):
        st.dataframe(result['diagnostics'], hide_index=True, use_container_width=True)
    with st.expander('ページ判定診断（本文文字数・キーワード一致）'):
        st.caption('HTTP 200かつHTMLとして取得できたページのみ。HTML内のみの一致は掲載確認ではありません。')
        st.dataframe(result['page_diagnostics'], hide_index=True, use_container_width=True)
    if not result['errors'].empty:
        with st.expander('エラー一覧'):
            st.dataframe(result['errors'], hide_index=True, use_container_width=True)
    domain_export = io.BytesIO()
    with pd.ExcelWriter(domain_export, engine='openpyxl') as writer:
        pd.DataFrame({'サイト名': [st.session_state.get('site_names', {}).get(d, '') for d in domains],
                      'ドメイン': domains}).to_excel(writer, index=False, sheet_name='ドメイン一覧')
        ws = writer.sheets['ドメイン一覧']
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 45
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name='Meiryo UI', size=10, bold=(cell.row == 1))
    st.download_button('📥 サイト名付きドメイン一覧をダウンロード', domain_export.getvalue(),
                       file_name='ドメイン一覧_サイト名付き.xlsx',
                       mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', on_click='ignore')
    st.download_button('📥 結果Excelをダウンロード', result['excel'],
                       file_name=result['filename'],
                       mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                       on_click='ignore')
    if result['error_count']:
        st.warning('一部処理できなかったドメインがあります。エラー一覧を確認してください。')
